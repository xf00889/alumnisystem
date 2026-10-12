import json
import zipfile
from datetime import datetime, timedelta
from io import BytesIO
from unittest.mock import patch

from django.conf import settings
from django.contrib.auth import get_user_model
from django.contrib.auth.models import AnonymousUser
from django.contrib.contenttypes.models import ContentType
from django.template.loader import render_to_string
from django.test import RequestFactory, TestCase, override_settings
from django.urls import resolve, reverse
from django.utils import timezone

from alumni_directory.models import Alumni
from core.admin_views import _tracer_program_choices, _tracer_targeting_from_form
from core.context_processors import tracer_study_banner_context
from core.models.notifications import Notification
from surveys.models import ResponseAnswer, Survey, SurveyQuestion, SurveyResponse
from surveys.tracer_study import (
    ALUMNI_TITLE,
    EMPLOYER_TITLE,
    _filtered_alumni_response_rows,
    _get_active_survey,
    alumni_is_eligible,
    eligible_alumni_queryset,
    public_tracer_studies_queryset,
    survey_accepting_responses,
)


TRACER_TEST_MIDDLEWARE = [
    middleware
    for middleware in settings.MIDDLEWARE
    if middleware != "setup.middleware.SetupRequiredMiddleware"
]


@override_settings(MIDDLEWARE=TRACER_TEST_MIDDLEWARE)
class TracerEligibilityRemediationTests(TestCase):
    def setUp(self):
        User = get_user_model()
        self.admin = User.objects.create_user(
            "tracer-admin", "admin@example.com", "pass", is_staff=True
        )
        self.match = self._alumni(
            "match", campus="BSC", college="CAS", course="BSCS", year=2001
        )
        self.wrong_campus = self._alumni(
            "campus", campus="MAIN", college="CAS", course="BSCS", year=2001
        )
        self.wrong_college = self._alumni(
            "college", campus="BSC", college="CTE", course="BSED", year=2001
        )
        self.wrong_program = self._alumni(
            "program", campus="BSC", college="CAS", course="BSINT", year=2001
        )
        self.wrong_year = self._alumni(
            "year", campus="BSC", college="CAS", course="BSCS", year=2004
        )

    def _alumni(self, username, *, campus, college, course, year):
        user = get_user_model().objects.create_user(
            username, f"{username}@example.com", "pass",
            first_name=username.title(), last_name="Alumni",
        )
        user.profile.has_completed_registration = True
        user.profile.save(update_fields=["has_completed_registration"])
        return Alumni.objects.create(
            user=user,
            campus=campus,
            college=college,
            course=course,
            graduation_year=year,
            gender="M",
            province="Negros Oriental",
            city="Dumaguete",
            address="Address",
        )

    def _survey(self, **overrides):
        values = {
            "title": ALUMNI_TITLE,
            "description": "SY 2026-2027 — Alumni",
            "created_by": self.admin,
            "start_date": timezone.now() - timedelta(days=1),
            "end_date": timezone.now() + timedelta(days=7),
            "status": "active",
            "display_to_all": False,
            "target_campus": "BSC",
            "target_college": "CAS",
            "target_program": "BSCS",
            "target_graduation_year_from": 2000,
            "target_graduation_year_to": 2003,
        }
        values.update(overrides)
        return Survey.objects.create(**values)

    def test_all_selected_target_conditions_are_anded(self):
        survey = self._survey()
        self.assertEqual(list(eligible_alumni_queryset(survey)), [self.match])
        self.assertTrue(alumni_is_eligible(survey, self.match))
        for alumnus in (
            self.wrong_campus,
            self.wrong_college,
            self.wrong_program,
            self.wrong_year,
        ):
            self.assertFalse(alumni_is_eligible(survey, alumnus))

    def test_restricted_without_conditions_fails_closed(self):
        survey = self._survey(
            target_campus="",
            target_college="",
            target_program="",
            target_graduation_year_from=None,
            target_graduation_year_to=None,
        )
        self.assertFalse(eligible_alumni_queryset(survey).exists())

    def test_newest_matching_restricted_cycle_wins_over_all_alumni_cycle(self):
        all_cycle = self._survey(
            description="All",
            display_to_all=True,
            target_campus="",
            target_college="",
            target_program="",
            target_graduation_year_from=None,
            target_graduation_year_to=None,
        )
        restricted = self._survey(description="Restricted")
        self.assertEqual(_get_active_survey(ALUMNI_TITLE, self.match), restricted)
        self.assertEqual(_get_active_survey(ALUMNI_TITLE, self.wrong_year), all_cycle)

    def test_direct_get_and_post_are_unavailable_to_out_of_scope_alumni(self):
        survey = self._survey()
        self.client.force_login(self.wrong_year.user)
        response = self.client.get(reverse("surveys:tracer_study_alumni"))
        self.assertRedirects(response, reverse("core:home"))
        response = self.client.post(reverse("surveys:tracer_study_alumni"), {})
        self.assertRedirects(response, reverse("core:home"))
        generic_url = reverse("surveys:survey_take", args=[survey.id])
        self.assertEqual(self.client.get(generic_url).status_code, 404)
        self.assertEqual(self.client.post(generic_url, {}).status_code, 404)
        self.assertFalse(SurveyResponse.objects.filter(survey=survey).exists())

        self.client.force_login(self.match.user)
        self.assertEqual(self.client.get(reverse("surveys:tracer_study_alumni")).status_code, 200)
        self.assertEqual(self.client.post(reverse("surveys:tracer_study_alumni"), {}).status_code, 200)
        self.assertTrue(SurveyResponse.objects.filter(survey=survey, alumni=self.match).exists())

    def test_primary_navigation_is_visible_only_to_authenticated_alumni(self):
        def render_nav(user):
            request = RequestFactory().get("/surveys/tracer-study/employer/")
            request.user = user
            request.resolver_match = resolve("/surveys/tracer-study/employer/")
            return render_to_string(
                "base.html",
                {"user": user, "request": request, "seo": {"title": "Tracer Study"}},
                request=request,
            )

        public_html = render_nav(AnonymousUser())
        self.assertNotIn('id="communityDropdown"', public_html)
        self.assertNotIn('id="careerDropdown"', public_html)

        alumni_html = render_nav(self.match.user)
        self.assertIn('id="communityDropdown"', alumni_html)
        self.assertIn('id="careerDropdown"', alumni_html)

    def test_alumni_tracer_form_uses_the_public_university_header(self):
        self._survey(show_on_public_page=True)
        self.client.force_login(self.match.user)
        response = self.client.get(reverse("surveys:tracer_study_alumni"))
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'class="university-header"')
        self.assertContains(response, 'class="navbar navbar-expand-lg landing-navbar"')
        self.assertNotContains(response, 'id="communityDropdown"')

    def test_banner_and_activation_notifications_use_the_same_cohort(self):
        survey = self._survey()
        content_type = ContentType.objects.get_for_model(Survey)
        recipients = set(Notification.objects.filter(
            content_type=content_type,
            object_id=survey.pk,
            notification_type="survey",
        ).values_list("recipient_id", flat=True))
        self.assertEqual(recipients, {self.match.user_id})

        class Resolver:
            url_name = "home"

        request = type("Request", (), {
            "user": self.match.user,
            "resolver_match": Resolver(),
        })()
        self.assertTrue(tracer_study_banner_context(request)["show_tracer_study_banner"])
        request.user = self.wrong_program.user
        self.assertFalse(tracer_study_banner_context(request)["show_tracer_study_banner"])

        survey.target_program = "BSINT"
        survey.save(update_fields=["target_program"])
        recipients = set(Notification.objects.filter(
            content_type=content_type,
            object_id=survey.pk,
            notification_type="survey",
        ).values_list("recipient_id", flat=True))
        self.assertEqual(recipients, {self.wrong_program.user_id})

    def test_period_boundaries_and_effective_states_are_inclusive(self):
        start = datetime(2026, 8, 15, 0, 0, 0)
        end = datetime(2026, 8, 20, 23, 59, 59, 999999)
        survey = self._survey(start_date=start, end_date=end)
        self.assertTrue(survey_accepting_responses(survey, start))
        self.assertTrue(survey_accepting_responses(survey, end))
        self.assertFalse(survey_accepting_responses(survey, start - timedelta(microseconds=1)))
        self.assertFalse(survey_accepting_responses(survey, end + timedelta(microseconds=1)))

        with patch("surveys.models.timezone.now", return_value=start - timedelta(days=1)):
            self.assertEqual(survey.availability_state(), "Scheduled")
        with patch("surveys.models.timezone.now", return_value=start):
            self.assertEqual(survey.availability_state(), "Open")
        with patch("surveys.models.timezone.now", return_value=end + timedelta(days=1)):
            self.assertEqual(survey.availability_state(), "Expired")
        survey.status = "closed"
        self.assertEqual(survey.availability_state(), "Manually Closed")

    def test_admin_targeting_validation_and_program_registry(self):
        with self.assertRaisesMessage(Exception, "at least one audience restriction"):
            _tracer_targeting_from_form({"visibility": "restricted"})
        with self.assertRaisesMessage(Exception, "before selecting a program"):
            _tracer_targeting_from_form({
                "visibility": "restricted", "program": "BSCS"
            })
        targeting = _tracer_targeting_from_form({
            "visibility": "restricted",
            "campus": "BSC",
            "college": "CAS",
            "program": "BSCS",
            "graduation_year_from": "2000",
            "graduation_year_to": "2003",
        })
        self.assertFalse(targeting["display_to_all"])
        self.assertEqual(targeting["target_program"], "BSCS")
        self.assertIn(
            ("CAS", "BSCS"),
            {(item["college"], item["code"]) for item in _tracer_program_choices()},
        )

    def test_admin_cycle_form_persists_combined_target_and_inclusive_end_date(self):
        self.client.force_login(self.admin)
        page = self.client.get(reverse("core:create_tracer_study"))
        self.assertContains(page, 'name="campus"')
        self.assertContains(page, 'name="graduation_year_from"')
        self.assertContains(page, 'data-college="CAS"')

        response = self.client.post(reverse("core:create_tracer_study"), {
            "audience": "alumni",
            "label": "SY 2026-2027",
            "start_date": "2026-08-15",
            "end_date": "2026-08-20",
            "visibility": "restricted",
            "campus": "BSC",
            "college": "CAS",
            "program": "BSCS",
            "graduation_year_from": "2000",
            "graduation_year_to": "2003",
        })
        self.assertRedirects(response, reverse("surveys:tracer_study_reports"))
        survey = Survey.objects.filter(title=ALUMNI_TITLE).latest("created_at")
        self.assertFalse(survey.display_to_all)
        self.assertEqual(
            (
                survey.target_campus,
                survey.target_college,
                survey.target_program,
                survey.target_graduation_year_from,
                survey.target_graduation_year_to,
            ),
            ("BSC", "CAS", "BSCS", 2000, 2003),
        )
        self.assertEqual(survey.end_date.time(), datetime.max.time())
        self.assertFalse(Survey.objects.filter(title=EMPLOYER_TITLE).exists())

    def test_admin_can_create_employer_questionnaire_without_alumni_targets(self):
        self.client.force_login(self.admin)
        page = self.client.get(reverse("core:create_tracer_study"))
        self.assertContains(page, 'name="audience"')
        self.assertContains(page, 'value="employer"')
        self.assertContains(page, 'id="employer-settings"')

        response = self.client.post(reverse("core:create_tracer_study"), {
            "audience": "employer",
            "label": "Employer Cycle 2026",
            "start_date": "2026-08-15",
            "end_date": "2026-08-20",
            # These must be ignored for the public employer questionnaire.
            "visibility": "restricted",
            "campus": "BSC",
            "college": "CAS",
            "program": "BSCS",
            "graduation_year_from": "2000",
            "graduation_year_to": "2003",
        })
        self.assertRedirects(response, reverse("surveys:tracer_study_reports"))
        survey = Survey.objects.get(title=EMPLOYER_TITLE)
        self.assertTrue(survey.display_to_all)
        self.assertEqual(survey.target_campus, "")
        self.assertEqual(survey.target_college, "")
        self.assertEqual(survey.target_program, "")
        self.assertIsNone(survey.target_graduation_year_from)
        self.assertIsNone(survey.target_graduation_year_to)
        self.assertTrue(survey.questions.exists())
        self.assertFalse(Survey.objects.filter(title=ALUMNI_TITLE).exists())


@override_settings(MIDDLEWARE=TRACER_TEST_MIDDLEWARE)
class PublicTracerStudyDiscoveryTests(TestCase):
    def setUp(self):
        self.admin = get_user_model().objects.create_user(
            "public-tracer-admin",
            "public-tracer-admin@example.com",
            "pass",
            is_staff=True,
        )
        self.now = timezone.now()

    def _survey(self, title, description, **overrides):
        values = {
            "title": title,
            "description": description,
            "created_by": self.admin,
            "start_date": self.now - timedelta(days=1),
            "end_date": self.now + timedelta(days=7),
            "status": "active",
            "display_to_all": True,
        }
        values.update(overrides)
        return Survey.objects.create(**values)

    def test_public_index_includes_all_employer_and_only_opted_in_alumni(self):
        employer = self._survey(
            EMPLOYER_TITLE,
            "Employer 2026 — Employer questionnaire",
            show_on_public_page=False,
        )
        visible_alumni = self._survey(
            ALUMNI_TITLE,
            "Alumni 2026 — Alumni questionnaire",
            show_on_public_page=True,
        )
        hidden_alumni = self._survey(
            ALUMNI_TITLE,
            "Hidden alumni — Alumni questionnaire",
            show_on_public_page=False,
        )

        self.assertEqual(
            set(public_tracer_studies_queryset()),
            {employer, visible_alumni},
        )
        response = self.client.get(reverse("surveys:tracer_study_public_list"))
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Employer 2026")
        self.assertContains(response, "Alumni 2026")
        self.assertNotContains(response, "Hidden alumni")
        self.assertContains(response, "Answer as Employer")
        self.assertContains(response, "Sign In to Respond")

    def test_anonymous_public_nav_badge_links_to_index_and_uses_index_anchor(self):
        self._survey(
            EMPLOYER_TITLE,
            "Employer 2026 — Employer questionnaire",
        )
        request = RequestFactory().get("/")
        request.user = AnonymousUser()
        request.resolver_match = resolve("/")
        context = tracer_study_banner_context(request)
        self.assertTrue(context["show_public_tracer_study_badge"])
        self.assertEqual(context["public_tracer_study_count"], 1)

        page = self.client.get(reverse("core:home"))
        self.assertContains(page, reverse("surveys:tracer_study_public_list"))
        self.assertContains(page, 'class="tracer-nav-badge"')
        self.assertNotContains(page, 'class="public-tracer-banner')
        index = self.client.get(reverse("surveys:tracer_study_public_list"))
        self.assertTrue(index.context["show_public_tracer_study_badge"])
        self.assertContains(index, 'class="tracer-nav-badge"')
        self.assertContains(index, 'id="available-tracer-studies"')
        self.assertNotContains(index, 'class="public-tracer-banner')

    def test_public_employer_form_uses_the_landing_page_header(self):
        self._survey(
            EMPLOYER_TITLE,
            "Employer 2026 — Employer questionnaire",
        )
        response = self.client.get(reverse("surveys:tracer_study_employer"))
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'class="university-header"')
        self.assertContains(response, 'class="navbar navbar-expand-lg landing-navbar"')
        self.assertNotContains(response, 'id="communityDropdown"')

    def test_create_form_controls_alumni_listing_but_employer_is_always_public(self):
        self.client.force_login(self.admin)
        create_url = reverse("core:create_tracer_study")
        page = self.client.get(create_url)
        self.assertContains(page, 'name="show_on_public_page"')

        common = {
            "start_date": (self.now - timedelta(days=1)).date().isoformat(),
            "end_date": (self.now + timedelta(days=7)).date().isoformat(),
            "visibility": "all",
        }
        response = self.client.post(
            create_url,
            {
                **common,
                "audience": "alumni",
                "label": "Public alumni cycle",
                "show_on_public_page": "on",
            },
        )
        self.assertRedirects(response, reverse("surveys:tracer_study_reports"))
        alumni = Survey.objects.get(title=ALUMNI_TITLE)
        self.assertTrue(alumni.show_on_public_page)

        response = self.client.post(
            create_url,
            {
                **common,
                "audience": "employer",
                "label": "Employer cycle",
            },
        )
        self.assertRedirects(response, reverse("surveys:tracer_study_reports"))
        employer = Survey.objects.get(title=EMPLOYER_TITLE)
        self.assertTrue(employer.show_on_public_page)

        edit_page = self.client.get(
            reverse("core:edit_tracer_study", args=[employer.id])
        )
        self.assertNotContains(edit_page, 'name="show_on_public_page"')


@override_settings(MIDDLEWARE=TRACER_TEST_MIDDLEWARE)
class TracerReportCohortRemediationTests(TestCase):
    def setUp(self):
        User = get_user_model()
        self.admin = User.objects.create_user(
            "report-admin", "report-admin@example.com", "pass", is_staff=True
        )
        self.client.force_login(self.admin)
        self.survey = Survey.objects.create(
            title=ALUMNI_TITLE,
            description="Report cohort",
            created_by=self.admin,
            start_date=timezone.now() - timedelta(days=1),
            end_date=timezone.now() + timedelta(days=7),
            status="active",
            display_to_all=False,
            target_college="CAS",
            target_program="BSCS",
            target_graduation_year_from=2000,
            target_graduation_year_to=2003,
        )
        self.question = SurveyQuestion.objects.create(
            survey=self.survey,
            question_text="Employer name",
            question_type="text",
            display_order=0,
            help_text=json.dumps({"part": "Employment"}),
        )
        self.year_2001 = self._alumni("Year One", 2001)
        self.year_2002 = self._alumni("Year Two", 2002)
        self.year_2003 = self._alumni("Year Three", 2003)
        self.out_scope = self._alumni("Out Scope", 2004)
        for alumnus, answer in (
            (self.year_2001, "Same employer"),
            (self.year_2002, "Same employer"),
            (self.out_scope, "Excluded employer"),
        ):
            response = SurveyResponse.objects.create(survey=self.survey, alumni=alumnus)
            ResponseAnswer.objects.create(
                response=response, question=self.question, text_answer=answer
            )

    def _alumni(self, name, year):
        slug = name.lower().replace(" ", "-")
        user = get_user_model().objects.create_user(
            slug, f"{slug}@example.com", "pass",
            first_name=name.split()[0], last_name=name.split()[1],
        )
        return Alumni.objects.create(
            user=user,
            campus="BSC",
            college="CAS",
            course="BSCS",
            graduation_year=year,
            gender="M",
            province="Negros Oriental",
            city="Dumaguete",
            address="Address",
        )

    def test_page_and_exports_share_year_filtered_official_cohort(self):
        params = {"year_from": "2002", "year_to": "2003"}
        page = self.client.get(
            reverse("surveys:tracer_study_report", args=[self.survey.id]), params
        )
        self.assertEqual(page.status_code, 200)
        self.assertEqual(page.context["total_responses"], 1)
        self.assertEqual(page.context["total_expected"], 2)
        self.assertEqual(len(page.context["responded_rows"]), 1)
        self.assertEqual(len(page.context["missing_rows"]), 1)
        self.assertEqual(len(page.context["out_of_scope_rows"]), 1)
        self.assertContains(page, "Out Scope")
        self.assertNotContains(page, "Excluded employer")

        export_url = reverse(
            "surveys:tracer_study_report_export", args=[self.survey.id]
        )
        csv_response = self.client.get(export_url, {**params, "format": "csv"})
        csv_text = csv_response.content.decode()
        self.assertIn("Year Two", csv_text)
        self.assertIn("Year Three", csv_text)
        self.assertNotIn("Year One", csv_text)
        self.assertNotIn("Out Scope", csv_text)

        excel_response = self.client.get(export_url, params)
        from openpyxl import load_workbook

        workbook = load_workbook(BytesIO(excel_response.content))
        self.assertEqual(workbook.sheetnames, ["Responded", "No Response", "Summary"])
        responded = [cell.value for row in workbook["Responded"].iter_rows() for cell in row]
        missing = [cell.value for row in workbook["No Response"].iter_rows() for cell in row]
        summary = workbook["Summary"]
        summary_values = [cell.value for row in summary.iter_rows() for cell in row]
        self.assertIn("Year Two", responded)
        self.assertNotIn("Year One", responded)
        self.assertIn("Year Three", missing)
        self.assertTrue(any(str(value).startswith("Q1:") for value in summary_values if value))
        employer_cell = next(
            cell for row in summary.iter_rows() for cell in row
            if cell.value == "Same employer"
        )
        self.assertEqual(summary.cell(employer_cell.row, 2).value, 1)
        self.assertIsInstance(summary.cell(employer_cell.row, 2).value, int)

    def test_zip_excludes_historical_out_of_scope_responses(self):
        class Driver:
            def get(self, url):
                pass

            def execute_cdp_cmd(self, name, params):
                if name == "Page.printToPDF":
                    import base64
                    return {"data": base64.b64encode(b"%PDF-1.4").decode("ascii")}
                return {}

            def quit(self):
                pass

        with patch("surveys.tracer_study._tracer_browser_driver", return_value=Driver()):
            response = self.client.get(
                reverse("surveys:tracer_study_report_export", args=[self.survey.id]),
                {"format": "zip"},
            )
        with zipfile.ZipFile(BytesIO(response.content)) as archive:
            names = "\n".join(archive.namelist())
        self.assertIn("One_Year", names)
        self.assertIn("Two_Year", names)
        self.assertNotIn("Scope_Out", names)
