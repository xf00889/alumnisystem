"""Staff-only employability statistics from submitted NORSU alumni tracer answers.

Quarterly hires require an actual current-job start date provided by the alumnus.
Submission date is never substituted for employment date.  Missing dates remain
unknown; alumni who never answered a tracer survey do not enter the respondent
denominator. Academic campus/college come from the related Alumni record, with
an explicit tracer-campus answer used when it maps unambiguously.
"""
import csv
import json
from collections import defaultdict
from datetime import date, datetime, timedelta, time
import calendar
from zoneinfo import ZoneInfo
from io import BytesIO

from django.conf import settings
from django.contrib.auth.decorators import login_required
from django.http import Http404, HttpResponse
from django.shortcuts import render
from django.views.decorators.http import require_GET

from alumni_directory.models import Alumni
from core.report_periods import export_download_name, parse_report_period
from .models import SurveyResponse
from .tracer_metadata import ALUMNI_TITLE
from .tracer_study import _can_view_tracer_reports


def _int_param(params, key, default, low=1900, high=2200):
    try:
        value = int(params.get(key) or default)
    except (TypeError, ValueError):
        raise ValueError(f'Choose a valid {key.replace("_", " ")}.') from None
    if not low <= value <= high:
        raise ValueError(f'Choose a valid {key.replace("_", " ")}.')
    return value


# Historical questionnaire campus options differ from alumni model codes.
# Do not guess which Bais campus was attended: use the verified profile for it.
TRACER_CAMPUS_CODES = {
    'main campus i': 'MAIN', 'main campus ii': 'MAIN',
    'bayawan-sta. catalina campus': 'BSC', 'siaton campus': 'SIATON',
    'guihulngan campus': 'GUI', 'mabinay campus': 'MAB',
    'pamplona campus': 'PAM',
}


def _answer_map(response):
    """Extract canonical questionnaire keys without using free-text guessing."""
    result = {}
    for answer in response.answers.all():
        question = answer.question
        try:
            meta = json.loads(question.help_text or '{}')
        except (TypeError, ValueError):
            continue
        key = meta.get('key')
        if key not in ('p2_year', 'p2_campus', 'p3_employed', 'p3_current_job_start'):
            continue
        if answer.selected_option_id:
            result[key] = answer.selected_option.option_text.strip()
        elif answer.text_answer:
            result[key] = answer.text_answer.strip()
    return result


def _parse_job_start(raw, as_of):
    try:
        value = date.fromisoformat(raw)
    except (ValueError, TypeError):
        return None
    return value if value <= as_of else None


def _period_buckets(period):
    """One selected bucket except annual reports, which show all four quarters.

    The labels, chart and export columns all derive from these exact windows.
    """
    if period.kind == 'annual':
        year = period.start.year
        return [
            {'label': f'Q{quarter} {year}',
             'start': date(year, (quarter - 1) * 3 + 1, 1),
             'end': date(year, quarter * 3, calendar.monthrange(year, quarter * 3)[1])}
            for quarter in range(1, 5)
        ]
    return [{'label': period.label, 'start': period.start, 'end': period.end}]


def _aggregate(cohort, hires, campus_names, college_names, employed_ids=None,
               undated_ids=None, buckets=None):
    """Aggregate distinct tracer respondents, counting starts only within display buckets."""
    if buckets is None:  # Preserve the helper's four-quarter behavior for existing callers.
        hire_year = next((started.year for _, started in hires), date.today().year)
        buckets = _period_buckets(type('Annual', (), {'kind': 'annual', 'start': date(hire_year, 1, 1)})())
    members = {}
    groups = defaultdict(lambda: {'graduates': 0, 'hires': [0] * len(buckets),
                                  'employed': 0, 'undated': 0})
    for pk, campus, college in cohort:
        if pk in members:
            continue
        group = (campus, college if campus == 'MAIN' else '')
        members[pk] = group
        groups[group]['graduates'] += 1
    first_hires = {}
    for alumni_id, started in hires:
        if alumni_id in members and (alumni_id not in first_hires or started < first_hires[alumni_id]):
            first_hires[alumni_id] = started
    if employed_ids is None:
        employed_ids = set(first_hires)
    employed_ids = set(employed_ids) & members.keys()
    undated_ids = set(undated_ids or ()) & employed_ids
    for alumni_id in employed_ids:
        groups[members[alumni_id]]['employed'] += 1
    for alumni_id in undated_ids:
        groups[members[alumni_id]]['undated'] += 1
    for alumni_id, started in first_hires.items():
        for index, bucket in enumerate(buckets):
            if bucket['start'] <= started <= bucket['end']:
                groups[members[alumni_id]]['hires'][index] += 1
                break

    def cells(counts, population):
        return [
            {'label': bucket['label'], 'count': count,
             'rate': round(100 * count / population, 2) if population else 0}
            for bucket, count in zip(buckets, counts)
        ]

    rows = []
    for (campus, college), counts in sorted(groups.items(), key=lambda x: (x[0][0] != 'MAIN', x[0][0], x[0][1])):
        population = counts['graduates']
        starts = counts['hires']
        rows.append({
            'campus': campus, 'college': college,
            'name': college_names.get(college, college or 'Unspecified') if campus == 'MAIN' else campus_names.get(campus, campus or 'Unspecified'),
            'graduates': population, 'hires': starts,
            'rates': [cell['rate'] for cell in cells(starts, population)],
            'period_cells': cells(starts, population),
            'total_hired': sum(starts), 'employed': counts['employed'], 'undated': counts['undated'],
            'rate': round(100 * counts['employed'] / population, 2) if population else 0,
        })

    def subtotal(group_rows, label):
        people = sum(row['graduates'] for row in group_rows)
        starts = [sum(row['hires'][i] for row in group_rows) for i in range(len(buckets))]
        employed = sum(row['employed'] for row in group_rows)
        return {
            'name': label, 'graduates': people, 'hires': starts,
            'rates': [cell['rate'] for cell in cells(starts, people)],
            'period_cells': cells(starts, people),
            'total_hired': sum(starts), 'employed': employed,
            'undated': sum(row['undated'] for row in group_rows),
            'rate': round(100 * employed / people, 2) if people else 0,
            'subtotal': True,
        }

    mains = [row for row in rows if row['campus'] == 'MAIN']
    externals = [row for row in rows if row['campus'] != 'MAIN']
    display_rows = []
    if mains:
        display_rows.extend(mains)
        display_rows.append(subtotal(mains, 'Main campus subtotal'))
    if externals:
        display_rows.extend(externals)
        display_rows.append(subtotal(externals, 'External campuses subtotal'))

    respondent_count = len(members)
    period_totals = [sum(row['hires'][i] for row in rows) for i in range(len(buckets))]
    employed_total = len(employed_ids)
    for row in display_rows:
        row['university_share'] = round(100 * row['graduates'] / respondent_count, 2) if respondent_count else 0
    return {
        'rows': rows, 'display_rows': display_rows, 'graduates': respondent_count,
        'period_labels': [bucket['label'] for bucket in buckets],
        'period_columns': buckets,
        'period_totals': period_totals,
        'period_total_cells': cells(period_totals, respondent_count),
        'table_column_count': 6 + 2 * len(buckets),
        # Preserve existing annual report / regression-test context keys.
        'quarter_totals': period_totals,
        'total_hired': sum(period_totals), 'employed': employed_total,
        'undated': len(undated_ids),
        'not_recorded': respondent_count - sum(period_totals),
        'rate': round(100 * employed_total / respondent_count, 2) if respondent_count else 0,
        'rate_display': f'{100 * employed_total / respondent_count:.2f}%' if respondent_count else 'N/A',
        'respondent_share_display': '100.00%' if respondent_count else 'N/A',
        'total_rates': [cell['rate'] for cell in cells(period_totals, respondent_count)],
    }


def _report_data(params):
    today = datetime.now(ZoneInfo(settings.TIME_ZONE)).date()
    cohort_year = _int_param(params, 'cohort_year', today.year - 1)
    report_year = _int_param(params, 'report_year', params.get('year') or today.year)
    raw_as_of = params.get('as_of') or today.isoformat()
    try:
        as_of = date.fromisoformat(raw_as_of)
    except (TypeError, ValueError):
        raise ValueError('Enter a valid as-of date.') from None
    if as_of > today:
        raise ValueError('As-of date cannot be in the future.')
    if report_year > as_of.year:
        raise ValueError('Report year cannot be later than the as-of year.')
    period_params = params if params.get('period') else {'period': 'annual', 'year': report_year}
    period = parse_report_period(period_params, today=today)
    # One-day reporting is intentionally unavailable for employability analytics.
    # The shared parser still supports it for unrelated administrative exports.
    # A same-day custom range is normalized to kind='day' by the parser.
    if period.kind == 'day':
        raise ValueError(
            'One-day employability reports are not available. Choose Monthly, '
            'Quarterly, Semi-annual, Annual, or a multi-day custom range.'
        )
    if not period.start:
        period = parse_report_period({'period':'annual', 'year':report_year}, today=today)
    if period.start.year != report_year or period.end.year != report_year:
        raise ValueError('Choose a timeframe within the selected hiring report year.')
    if period.start > as_of:
        raise ValueError('The selected reporting period starts after the as-of date.')
    campus_filter = (params.get('campus') or '').strip()
    valid_campuses = dict(Alumni.CAMPUS_CHOICES)
    if campus_filter and campus_filter not in valid_campuses:
        raise ValueError('Choose a valid campus.')

    # Only submitted alumni tracer questionnaires count. The latest response
    # as of the snapshot date wins when an alumnus answered multiple cycles.
    cutoff = datetime.combine(as_of + timedelta(days=1), time.min)
    response_query = (SurveyResponse.objects.filter(
        survey__title=ALUMNI_TITLE, submitted_at__lt=cutoff,
    ).select_related('alumni').prefetch_related('answers__question', 'answers__selected_option')
     .order_by('alumni_id', '-submitted_at', '-pk'))
    seen, cohort, hires, employed_ids, undated_ids = set(), [], [], set(), set()
    for response in response_query:
        if response.alumni_id in seen:
            continue
        seen.add(response.alumni_id)
        alumni = response.alumni
        answers = _answer_map(response)
        year = alumni.graduation_year
        try:
            year = int(answers.get('p2_year') or year)
        except (ValueError, TypeError):
            pass
        if year != cohort_year:
            continue
        campus = TRACER_CAMPUS_CODES.get((answers.get('p2_campus') or '').casefold(), alumni.campus)
        if campus_filter and campus != campus_filter:
            continue
        cohort.append((response.alumni_id, campus, alumni.college))
        if (answers.get('p3_employed') or '').strip().casefold() != 'yes':
            continue
        employed_ids.add(response.alumni_id)
        raw_start = answers.get('p3_current_job_start')
        started = _parse_job_start(raw_start, as_of) if raw_start else None
        if started is None:
            undated_ids.add(response.alumni_id)
        elif period.start <= started <= min(as_of, period.end):
            hires.append((response.alumni_id, started))

    result = _aggregate(cohort, hires, valid_campuses, dict(Alumni.COLLEGE_CHOICES),
                        employed_ids=employed_ids, undated_ids=undated_ids,
                        buckets=_period_buckets(period))
    result.update({
        'cohort_year': cohort_year, 'report_year': report_year, 'as_of': as_of,
        'campus_filter': campus_filter, 'campus_choices': Alumni.CAMPUS_CHOICES,
        'generation_date': today, 'period': period, 'period_kind': period.kind,
        'period_cutoff': min(as_of, period.end),
        'methodology': (
            'Source: submitted NORSU alumni tracer-study questionnaire responses, latest response per alumnus '
            'received on or before the as-of date. The cohort denominator is respondents, not the official '
            'graduate registry. Employment rate uses the answer to "Are you presently employed?". '
            'Quarterly counts require an explicit current-job start date, and only starts in the selected '
            'timeframe are counted. Missing or invalid dates are reported as undated, never assigned to '
            'the response-submission quarter. Campus may come from the tracer answer; college comes from '
            'the linked alumni record. These are self-reported responses, not independently verified hires.'),
    })
    return result


def _table_rows(report):
    for row in report['display_rows']:
        values = [row['name'], row['graduates']]
        for cell in row['period_cells']:
            values += [cell['count'], f"{cell['rate']:.2f}%"]
        yield values + [row['employed'], row['undated'],
                        f"{row['rate']:.2f}%", f"{row['university_share']:.2f}%"]
    totals = ['UNIVERSITY TOTAL', report['graduates']]
    for cell in report['period_total_cells']:
        totals += [cell['count'], f"{cell['rate']:.2f}%"]
    yield totals + [report['employed'], report['undated'],
                    report['rate_display'], report['respondent_share_display']]


def _periodized_export_report(report):
    """Keep exports compatible with existing directory reporting callers.

    Older directory code calls _aggregate without a period (four quarter
    buckets), then attaches a monthly/quarterly/half-year ReportPeriod.
    Those dated starts were already filtered to the selected window.
    Collapse their quarter bins to exactly the requested export window.
    """
    period = report['period']
    if period.kind == 'annual' or len(report['period_labels']) != 4:
        return report
    if period.kind not in ('monthly', 'quarterly', 'semi_annual', 'day', 'custom'):
        return report
    quarter_indices = set(range((period.start.month - 1) // 3,
                                (period.end.month - 1) // 3 + 1))
    def collapse(row):
        entry = dict(row)
        count = sum(row['hires'][i] for i in quarter_indices)
        denom = row['graduates']
        rate = round(100 * count / denom, 2) if denom else 0
        entry['period_cells'] = [{'label': period.label, 'count': count, 'rate': rate}]
        return entry
    result = dict(report)
    result['display_rows'] = [collapse(row) for row in report['display_rows']]
    result['period_labels'] = [period.label]
    total = sum(report['quarter_totals'][i] for i in quarter_indices)
    denom = report['graduates']
    result['period_total_cells'] = [{
        'label': period.label, 'count': total,
        'rate': round(100 * total / denom, 2) if denom else 0,
    }]
    result['period_totals'] = [total]
    result['total_hired'] = total
    result['rate_display'] = f"{report['rate']:.2f}%" if denom else 'N/A'
    return result


def _export_report(report, format_name):
    report = _periodized_export_report(report)
    period = report['period']
    base = f"employability_cohort-{report['cohort_year']}_as-of-{report['as_of'].isoformat()}"
    columns = ['Campus / College', 'Tracer respondents'] + [name for label in report['period_labels'] for name in (f'Starts {label}', f'Start rate {label}')] + ['Reported employed (as of)', 'Undated employed', 'Respondent employment rate (as of)', 'Share of respondents']
    rows = list(_table_rows(report))
    if format_name == 'csv':
        response = HttpResponse(content_type='text/csv; charset=utf-8')
        response['Content-Disposition'] = f'attachment; filename="{export_download_name(base, period, "csv")}"'
        writer = csv.writer(response)
        writer.writerow(['NORSU OFFICE OF ALUMNI AFFAIRS - TRACER-RESPONSE EMPLOYABILITY REPORT'])
        writer.writerow(['Graduating cohort', report['cohort_year'], 'Current job starts in year', report['report_year'], 'Period', period.label, 'As of', report['as_of']])
        writer.writerow([report['methodology']])
        writer.writerow(columns)
        writer.writerows(rows)
        return response
    if format_name in ('excel', 'xlsx'):
        from openpyxl import Workbook
        from openpyxl.styles import Font, Alignment, PatternFill
        from openpyxl.utils import get_column_letter
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = 'Employability'
        sheet.append(['NEGROS ORIENTAL STATE UNIVERSITY — OFFICE OF ALUMNI AFFAIRS'])
        sheet.append([f"GRADUATE EMPLOYABILITY REPORT | Class {report['cohort_year']} | {period.label} | As of {report['as_of']}"])
        sheet.append([report['methodology']])
        sheet.append([])
        sheet.append(columns)
        for values in rows:
            sheet.append(values)
        sheet.freeze_panes = 'C6'
        sheet.auto_filter.ref = f'A5:{get_column_letter(len(columns))}{sheet.max_row}'
        sheet.column_dimensions['A'].width = 44
        for column in range(2, len(columns) + 1):
            sheet.column_dimensions[get_column_letter(column)].width = 19
        for row in (1, 2):
            sheet.cell(row=row, column=1).font = Font(bold=True, color='172541', size=13)
        sheet.row_dimensions[3].height = 45
        sheet.cell(row=3, column=1).alignment = Alignment(wrap_text=True)
        for cell in sheet[5]:
            cell.fill = PatternFill('solid', fgColor='172541')
            cell.font = Font(bold=True, color='FFFFFF')
            cell.alignment = Alignment(wrap_text=True)
        for cell in sheet[sheet.max_row]:
            cell.font = Font(bold=True)
        output = BytesIO()
        workbook.save(output)
        response = HttpResponse(output.getvalue(), content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        response['Content-Disposition'] = f'attachment; filename="{export_download_name(base, period, "xlsx")}"'
        return response
    if format_name == 'pdf':
        from reportlab.lib import colors
        from reportlab.lib.enums import TA_CENTER
        from reportlab.lib.pagesizes import A4, landscape
        from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
        from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, KeepTogether
        output = BytesIO()
        document = SimpleDocTemplate(output, pagesize=landscape(A4), leftMargin=34, rightMargin=34, topMargin=28, bottomMargin=32)
        styles = getSampleStyleSheet()
        heading = ParagraphStyle('norsu-heading', parent=styles['Heading2'], textColor=colors.HexColor('#172541'), alignment=TA_CENTER)
        from pathlib import Path
        from django.conf import settings
        from reportlab.platypus import Image
        header_path = Path(settings.BASE_DIR) / 'static' / 'images' / 'norsu-header.png'
        elements = []
        if header_path.is_file():
            elements.append(Image(str(header_path), width=670, height=63))
        elements += [Paragraph('NEGROS ORIENTAL STATE UNIVERSITY', heading),
                    Paragraph('Office of Alumni Affairs - Tracer-response Employability Report', styles['Heading3']),
                    Paragraph(f"Graduating cohort {report['cohort_year']} | Reporting window {period.label} | As of {report['as_of']}", styles['Normal']),
                    Spacer(1, 12)]
        white_style = ParagraphStyle('norsu-col-header', parent=styles['Normal'], textColor=colors.white, fontSize=7)
        cell_style = ParagraphStyle('norsu-table-cell', parent=styles['Normal'], fontSize=8, leading=10)
        data = [[Paragraph(str(cell), white_style) for cell in columns]]
        for record in rows:
            data.append([Paragraph(str(cell), cell_style) for cell in record])
        # Adapt widths to one selected period or four annual quarter columns.
        available = landscape(A4)[0] - 68
        first_width = 146 if len(report['period_labels']) == 4 else 176
        respondent_width = 64 if len(report['period_labels']) == 4 else 73
        tail_widths = [60, 52, 62, 62] if len(report['period_labels']) == 4 else [86, 72, 86, 83]
        period_width = (available - first_width - respondent_width - sum(tail_widths)) / (2 * len(report['period_labels']))
        widths = [first_width, respondent_width] + [period_width] * (2 * len(report['period_labels'])) + tail_widths
        tbl = Table(data, colWidths=widths, repeatRows=1, hAlign='CENTER')
        tbl.setStyle(TableStyle([
            ('BACKGROUND', (0,0), (-1,0), colors.HexColor('#172541')),
            ('TEXTCOLOR', (0,0), (-1,0), colors.whitesmoke),
            ('ROWBACKGROUNDS', (0,1), (-1,-2), [colors.white, colors.HexColor('#F5F2EC')]),
            ('BACKGROUND', (0,-1), (-1,-1), colors.HexColor('#E7D8B9')),
            ('GRID', (0,0), (-1,-1), .4, colors.HexColor('#D5D7DB')),
            ('VALIGN', (0,0), (-1,-1), 'MIDDLE'),
            ('LEFTPADDING', (0,0), (-1,-1), 3), ('RIGHTPADDING', (0,0), (-1,-1), 3),
            ('TOPPADDING', (0,0), (-1,-1), 7), ('BOTTOMPADDING', (0,0), (-1,-1), 7),
        ]))
        elements += [tbl, Spacer(1, 10), Paragraph(report['methodology'], styles['Normal']),
                     Spacer(1, 12), Paragraph('Generated by NORSU Alumni Portal | Office of Alumni Affairs | Official registry reconciliation required', styles['Italic'])]
        document.build(elements)
        response = HttpResponse(output.getvalue(), content_type='application/pdf')
        response['Content-Disposition'] = f'attachment; filename="{export_download_name(base, period, "pdf")}"'
        return response
    raise Http404('Unsupported export format')


@login_required
@require_GET
def employability_report(request):
    if not _can_view_tracer_reports(request.user):
        raise Http404
    try:
        data = _report_data(request.GET)
    except ValueError as error:
        return render(request, 'tracer_study/employability_report.html', {
            'error': str(error), 'campus_choices': Alumni.CAMPUS_CHOICES,
            'cohort_year': request.GET.get('cohort_year', ''),
            'report_year': request.GET.get('report_year', ''),
            'as_of': request.GET.get('as_of', ''),
            'campus_filter': request.GET.get('campus', ''),
            'period_kind': (request.GET.get('period') if request.GET.get('period') in
                            ('monthly', 'quarterly', 'semi_annual', 'annual', 'custom') else 'annual'),
        }, status=400)
    format_name = (request.GET.get('format') or '').lower()
    if format_name:
        return _export_report(data, format_name)
    return render(request, 'tracer_study/employability_report.html', data)
